"""Surat pengantar (cover letter) untuk submission ke European Journal of Agronomy."""
import datetime as dt
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for a in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
    setattr(sec, a, Cm(2.5))
st = doc.styles["Normal"]
st.font.name = "Arial"; st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(8); st.paragraph_format.line_spacing = 1.15

def P(t, bold=False, italic=False, align=None):
    p = doc.add_paragraph(); r = p.add_run(t)
    r.bold = bold; r.italic = italic
    if align is not None: p.alignment = align
    return p

P(dt.date.today().strftime("%d %B %Y"), align=WD_ALIGN_PARAGRAPH.RIGHT)
P("Editor-in-Chief\nEuropean Journal of Agronomy")
P("Dear Editor,", bold=True)
P('Please consider our manuscript "Nitrogen-limited WOFOST 8.1 for tropical transplanted rice: '
  'parameter-consistency fixes, a leaf-level nitrogen extension and independent omission-plot validation '
  'in West Java, Indonesia" for publication as a research article in the European Journal of Agronomy.')
P("Nitrogen-limited WOFOST is increasingly used to guide fertiliser management in tropical rice, yet its "
  "recent 8.1 implementation has rarely been tested for internal consistency or against independent nitrogen-"
  "omission data. Our study addresses this gap with three contributions:")
P("1. We identify and correct two parameter-consistency problems in WOFOST 8.1 as distributed with PCSE 6.0.13 "
  "(a relative-growth-rate floor that inverts the nitrogen stress response for slow-growing tropical cultivars, "
  "and a reference assimilation rate that silently ignores the calibrated photosynthesis table), which are "
  "relevant to any user calibrating the model for transplanted rice.")
P("2. We show mechanistically, using tillering, SPAD and light-interception evidence, that standard WOFOST 8.1 "
  "dilutes leaf nitrogen at conserved leaf area whereas tropical rice conserves leaf nitrogen per unit area and "
  "reduces leaf area; a single-parameter leaf-level extension (LINTUL3-type) restores the observed canopy "
  "strategy and more than doubles the reliability of probabilistic LAI predictions.")
P("3. All calibrations are evaluated by staged Bayesian inference with full convergence diagnostics, "
  "leave-one-dose-out cross-validation and, importantly, blind prediction of two independent seasons of 0-N "
  "omission plots from a long-term fertility experiment, with the indigenous nitrogen supply reproduced within "
  "2%. The complete application (WOFOST Studio), all literature-mined data and every analysis script are openly "
  "archived (GitHub; Zenodo DOI 10.5281/zenodo.22969839), so every number in the manuscript can be regenerated.")
P("The manuscript is approximately 11,200 words with 8 figures, 5 tables and supplementary material. It has not "
  "been published previously, is not under consideration elsewhere, and all authors have approved the submission. "
  "The use of an AI assistant in software development and manuscript drafting is disclosed in the manuscript, and "
  "the authors take full responsibility for its content. We have no competing interests to declare.")
P("We suggest reviewers with expertise in WOFOST/PCSE nitrogen modelling, rice nitrogen physiology, or Bayesian "
  "calibration of crop models. [Optional: insert 3–5 named reviewers with e-mail addresses here.]", italic=True)
P("Thank you for considering our work.")
P("Yours sincerely,")
P("Zainal Arifin (corresponding author)\n"
  "Department of Agribusiness, Vocational School, Universitas Sebelas Maret, Indonesia\n"
  "zainal.arifin@staff.uns.ac.id\n"
  "on behalf of all authors, including Iwan Gunawan (Department of Mechanical Engineering, Universitas Khairun, "
  "Ternate 97719, Indonesia; iwan99gun@unkhair.ac.id)")

out = Path(__file__).with_name("surat_pengantar_EJA.docx")
doc.save(out)
print("DOCX:", out)
try:
    import win32com.client
    w = win32com.client.Dispatch("Word.Application"); w.Visible = False
    d = w.Documents.Open(str(out.resolve()))
    pdf = out.with_suffix(".pdf")
    d.SaveAs2(str(pdf.resolve()), FileFormat=17)
    d.Close(False); w.Quit()
    print("PDF:", pdf)
except Exception as e:
    print("Konversi PDF gagal:", e)
