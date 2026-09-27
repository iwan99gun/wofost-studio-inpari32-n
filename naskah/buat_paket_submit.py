"""Susun folder paket unggahan untuk European Journal of Agronomy (Editorial Manager) di
D:/riset_tani_1/SUBMISSION_EJA/ :
  01_Manuscript_EJA.docx          naskah bersih 1 kolom (sumber yang dapat diedit; wajib)
  02_Highlights.docx              3-5 poin, <= 85 karakter (wajib)
  03_Graphical_abstract.tif/.png  (wajib saat submit)
  04_Cover_letter.docx
  05_Declaration_of_interest.docx
  Figures/Figure_1..6.tif (+png), Figure_S1..S2.tif (+png)   - satu berkas per gambar
  DATA_FORMULIR.txt               teks siap-salin untuk kolom formulir (judul, abstrak, kata kunci, penulis, dll.)
  CHECKLIST.txt                   langkah unggah + hal yang harus diputuskan penulis
Jalankan setelah: python buat_naskah.py final ; python buat_surat.py ; python buat_graphical_abstract.py"""
import shutil, re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PKG = ROOT / "SUBMISSION_EJA"
if PKG.exists():
    shutil.rmtree(PKG)
(PKG / "Figures").mkdir(parents=True)

# ---- 1. naskah
shutil.copy(HERE / "Manuscript_EJA.docx", PKG / "01_Manuscript_EJA.docx")
src = Document(HERE / "Manuscript_EJA.docx")
paras = [p.text for p in src.paragraphs]

# ---- 2. highlights
i = paras.index("Highlights")
hl = []
for t in paras[i + 1:i + 8]:
    if t.startswith("• "):
        hl.append(t[2:].strip())
assert 3 <= len(hl) <= 5 and all(len(h) <= 85 for h in hl), [len(h) for h in hl]
d = Document(); d.styles["Normal"].font.name = "Arial"; d.styles["Normal"].font.size = Pt(11)
d.add_heading("Highlights", level=1)
for h in hl:
    d.add_paragraph(h, style="List Bullet")
d.save(PKG / "02_Highlights.docx")

# ---- 3. graphical abstract
for ext in ("tif", "png"):
    shutil.copy(HERE / "gambar" / f"Graphical_abstract.{ext}", PKG / f"03_Graphical_abstract.{ext}")

# ---- 4. cover letter
shutil.copy(HERE / "surat_pengantar_EJA.docx", PKG / "04_Cover_letter.docx")

# ---- 5. declaration of interest
d = Document(); d.styles["Normal"].font.name = "Arial"; d.styles["Normal"].font.size = Pt(11)
d.add_heading("Declaration of interests", level=1)
d.add_paragraph("\u2612 The authors declare that they have no known competing financial interests or personal relationships "
                "that could have appeared to influence the work reported in this paper.")
d.add_paragraph("\u2610 The authors declare the following financial interests/personal relationships which may be considered "
                "as potential competing interests:")
d.save(PKG / "05_Declaration_of_interest.docx")

# ---- 6. gambar (satu berkas per gambar, penamaan logis)
MAP = {"Fig1_workflow": "Figure_1", "Fig2_potential_calibration": "Figure_2", "Fig3_N_response": "Figure_3",
       "Fig4_NNI_LAI_mechanism": "Figure_4", "Fig5_independent_validation": "Figure_5", "Fig6_Sobol": "Figure_6",
       "FigS1_weather_source": "Figure_S1", "FigS2_variety_assumption": "Figure_S2"}
for k, v in MAP.items():
    for ext in ("tif", "png"):
        shutil.copy(HERE / "gambar" / f"{k}.{ext}", PKG / "Figures" / f"{v}.{ext}")

# ---- 7. data formulir
ia = paras.index("Abstract")
abstract = paras[ia + 1]
kw = next(t for t in paras if t.startswith("Keywords:"))[len("Keywords:"):].strip()
title = paras[0]
sw = next(t for t in paras if t.startswith("WOFOST Studio (Python"))
form = f"""DATA UNTUK FORMULIR EDITORIAL MANAGER (salin-tempel)
=====================================================
Tipe artikel      : Original research paper (Regular paper)
Judul             :
{title}

Penulis (URUTAN harus sama dengan halaman judul naskah):
  1. Zainal Arifin  (PENULIS KORESPONDENSI)
     Department of Agribusiness, Vocational School, Universitas Sebelas Maret, Indonesia
     zainal.arifin@staff.uns.ac.id   ORCID 0009-0008-0345-3167
  2. Iwan Gunawan
     Department of Mechanical Engineering, Universitas Khairun, Ternate 97719, Indonesia
     iwan99gun@unkhair.ac.id         ORCID 0000-0002-2784-4436

Abstrak ({len(abstract.split())} kata):
{abstract}

Kata kunci (Elsevier: 1-7 kata kunci, dipisah titik koma):
{kw}

Pernyataan ketersediaan data / kode ("Data statement"):
{sw}

Pernyataan konflik kepentingan: penulis tidak memiliki konflik kepentingan (lihat 05_Declaration_of_interest.docx).
Pernyataan penggunaan AI generatif: sudah ada di naskah (bagian sebelum Referensi).
Pendanaan: LIHAT CHECKLIST.txt (belum ditulis; harus Anda isi).
"""
(PKG / "DATA_FORMULIR.txt").write_text(form, encoding="utf-8")

# ---- 8. checklist
chk = """CHECKLIST UNGGAH - European Journal of Agronomy  (https://submit.elsevier.com/EURAGR)
=====================================================================================
YANG HARUS ANDA PUTUSKAN / ISI SEBELUM UNGGAH (saya tidak boleh menebaknya)
 [ ] PENDANAAN. Panduan Elsevier meminta pernyataan sumber dana. Naskah belum memuatnya. Kalau tidak ada dana khusus,
     Elsevier menyarankan kalimat: "This research did not receive any specific grant from funding agencies in the
     public, commercial, or not-for-profit sectors." Isi sesuai kenyataan, lalu minta saya menambahkannya ke naskah.
 [ ] ALAMAT POS AFILIASI. Panduan meminta nama lengkap dan alamat pos setiap afiliasi, termasuk negara.
     Afiliasi Zainal Arifin sekarang hanya "...Universitas Sebelas Maret, Indonesia" (kota/kode pos belum ada).
     Afiliasi Iwan Gunawan sudah punya kota dan kode pos (Ternate 97719).
 [ ] PENULIS KORESPONDENSI = orang yang login dan mengunggah. Sekarang naskah menyebut Zainal Arifin. Kalau yang
     mengunggah Iwan Gunawan, minta saya ubah halaman judul + surat pengantar dulu. Tidak bisa diganti setelah diterima.
 [ ] Nama calon reviewer (opsional) di 04_Cover_letter.docx; editor biasanya meminta 3-5 nama + surel di formulir.
 [ ] Kedua penulis sudah membaca dan menyetujui naskah final.

URUTAN UNGGAH DI SISTEM (jenis berkas di menu "Attach Files")
 1. Manuscript          -> 01_Manuscript_EJA.docx
 2. Highlights          -> 02_Highlights.docx
 3. Graphical Abstract  -> 03_Graphical_abstract.tif  (atau .png)
 4. Figure (satu per satu, urut) -> Figures/Figure_1.tif ... Figure_6.tif   (keterangan gambar ada di dalam naskah)
 5. Supplementary Material -> Figure_S1.tif, Figure_S2.tif  (tabel S1-S5 sudah ada di akhir naskah)
 6. Declaration of Interest -> 05_Declaration_of_interest.docx
 7. Cover Letter        -> 04_Cover_letter.docx
 Setelah semua terunggah: klik "Build PDF for my Approval", PERIKSA PDF-nya (gambar, tabel, nomor baris), baru "Approve".

PERIKSA CEPAT
 - Naskah: 1 kolom, nomor baris, tabel sebagai teks yang bisa diedit (sudah).
 - Gambar: TIFF 600 dpi, lebar 190 mm (memenuhi syarat resolusi Elsevier).
 - Semua referensi dalam teks ada di daftar dan sebaliknya (sudah diperiksa otomatis).
 - Data & kode: https://github.com/iwan99gun/wofost-studio-inpari32-n  |  DOI https://doi.org/10.5281/zenodo.22969839
   (rekam Zenodo baru mencantumkan Iwan Gunawan; tambahkan Zainal Arifin lewat tombol Edit bila perlu.)
"""
(PKG / "CHECKLIST.txt").write_text(chk, encoding="utf-8")

# ---- ringkasan
for p in sorted(PKG.rglob("*")):
    if p.is_file():
        print(f"{p.relative_to(PKG)}   {p.stat().st_size / 1024:.0f} KB")
print("paket:", PKG)
